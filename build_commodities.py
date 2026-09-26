r"""build_commodities.py - fuel & equipment prices (roadmap II.C 'Fuel prices / equipment cost trends') -> data\commodities.json

INPUTS (desktop downloads in data\CommodityPricing\, 2026-09-26):
  MHHNGSP.xlsx          FRED - Henry Hub natural gas spot, $/MMBtu, monthly
  PURANUSDM.xlsx        FRED - global price of uranium, $/lb, monthly (IMF benchmark)
  PCU335311335311.xlsx  FRED - PPI, electric power & specialty transformer manufacturing, index Jun-1981=100, monthly
  Coal_shipments_to_the_electric_power_sector__price_by_plant_state.csv
                        EIA coal data browser - delivered coal price to the power sector by plant state, $/short ton, quarterly
  Market_average_price.csv
                        EIA coal data browser - market average (mine) price by basin, open market vs captive, $/short ton, annual
  (Coal SPOT history is proprietary to S&P and is not published by EIA; delivered and mine prices are the public series.)

PER-NAME EXPOSURE (off data\fleet.json, operated view): gas and coal share of operable MW, and the coal-MW-weighted
delivered coal price across the states where the name burns coal (latest quarter vs a year earlier). Exposure is
descriptive: most regulated fuel cost passes through fuel clauses; merchant names (VST, TLN) carry it in margin.

    python scripts\build_commodities.py [data_dir]
Re-run after refreshing any of the five downloads (same file names).
"""
import sys, os, csv, json, datetime, statistics, collections
import openpyxl

DATA = sys.argv[1] if len(sys.argv) > 1 else r'E:\PowerAcademy\data'
SRC = os.path.join(DATA, 'CommodityPricing')
STATE_ABBR = {'Alabama': 'AL', 'Alaska': 'AK', 'Arizona': 'AZ', 'Arkansas': 'AR', 'California': 'CA', 'Colorado': 'CO',
    'Connecticut': 'CT', 'Delaware': 'DE', 'Florida': 'FL', 'Georgia': 'GA', 'Hawaii': 'HI', 'Idaho': 'ID', 'Illinois': 'IL',
    'Indiana': 'IN', 'Iowa': 'IA', 'Kansas': 'KS', 'Kentucky': 'KY', 'Louisiana': 'LA', 'Maine': 'ME', 'Maryland': 'MD',
    'Massachusetts': 'MA', 'Michigan': 'MI', 'Minnesota': 'MN', 'Mississippi': 'MS', 'Missouri': 'MO', 'Montana': 'MT',
    'Nebraska': 'NE', 'Nevada': 'NV', 'New Hampshire': 'NH', 'New Jersey': 'NJ', 'New Mexico': 'NM', 'New York': 'NY',
    'North Carolina': 'NC', 'North Dakota': 'ND', 'Ohio': 'OH', 'Oklahoma': 'OK', 'Oregon': 'OR', 'Pennsylvania': 'PA',
    'Rhode Island': 'RI', 'South Carolina': 'SC', 'South Dakota': 'SD', 'Tennessee': 'TN', 'Texas': 'TX', 'Utah': 'UT',
    'Vermont': 'VT', 'Virginia': 'VA', 'Washington': 'WA', 'West Virginia': 'WV', 'Wisconsin': 'WI', 'Wyoming': 'WY',
    'District of Columbia': 'DC'}


def fred(fn):
    wb = openpyxl.load_workbook(os.path.join(SRC, fn), read_only=True, data_only=True)
    ws = wb['Monthly'] if 'Monthly' in wb.sheetnames else wb.worksheets[-1]
    out, title = [], None
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], datetime.datetime) and isinstance(r[1], (int, float)):
            out.append((r[0].strftime('%Y-%m'), float(r[1])))
    rd = wb['README'] if 'README' in wb.sheetnames else None
    if rd:
        for r in rd.iter_rows(values_only=True):
            if r and r[0] and r[1] and str(r[0]).strip() == fn.split('.')[0]:
                title = str(r[1]); upd = str(r[2] or '')
    return out, title


def summarise(series, unit, since='2015-01'):
    s = [(d, v) for d, v in series if d >= since]
    last_d, last_v = series[-1]
    by_d = dict(series)
    y_ago = by_d.get('%d-%s' % (int(last_d[:4]) - 1, last_d[5:]))
    last12 = [v for d, v in series[-12:]]
    ann = collections.defaultdict(list)
    for d, v in series:
        ann[d[:4]].append(v)
    annual = {y: round(statistics.mean(v), 3) for y, v in sorted(ann.items()) if len(v) == 12 or y == last_d[:4]}
    five = [v for d, v in series if int(d[:4]) >= int(last_d[:4]) - 5 and d <= last_d][-60:]
    return {'unit': unit, 'latest': {'month': last_d, 'value': round(last_v, 3)},
            'yoy_pct': round(100 * (last_v / y_ago - 1), 1) if y_ago else None,
            'avg_12m': round(statistics.mean(last12), 3), 'avg_5y': round(statistics.mean(five), 3) if five else None,
            'min_since_2015': round(min(v for _, v in s), 3), 'max_since_2015': round(max(v for _, v in s), 3),
            'annual_avg': annual, 'monthly': [[d, round(v, 3)] for d, v in s]}


def eia_csv(fn):
    rows = list(csv.reader(open(os.path.join(SRC, fn), encoding='utf-8-sig')))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == 'description')
    cols = rows[hi][3:]
    out = collections.OrderedDict()
    for r in rows[hi + 1:]:
        if len(r) < 4 or not r[2]:
            continue
        vals = {}
        for c, v in zip(cols, r[3:]):
            try:
                vals[c] = float(v)
            except ValueError:
                pass
        if vals:
            out[r[0]] = {'units': r[1], 'key': r[2], 'values': vals}
    return cols, out, rows[1][0] if len(rows) > 1 else None


def main():
    doc = collections.OrderedDict()
    doc['_schema_version'] = '1.0'
    doc['_generated'] = datetime.date.today().isoformat()
    doc['_sources'] = {'MHHNGSP': 'FRED - Henry Hub natural gas spot price (EIA), $/MMBtu, monthly',
                       'PURANUSDM': 'FRED - Global price of uranium (IMF), $/lb, monthly',
                       'PCU335311335311': 'FRED - PPI by industry: electric power & specialty transformer manufacturing (BLS), Jun-1981=100',
                       'coal_delivered': 'EIA coal data browser - coal shipments to the electric power sector: price by plant state, $/short ton, quarterly',
                       'coal_mine': 'EIA coal data browser - market average price by basin (open market / captive), $/short ton, annual'}
    doc['_caveat'] = ('Public series only: coal SPOT history is S&P-proprietary, so coal is shown as the delivered price to power plants '
                      '(what utilities pay, incl. transport) and the mine-mouth market average by basin. Uranium is the IMF global '
                      'benchmark (not a term-contract price). The transformer PPI is an equipment-cost proxy, not a quote. '
                      'Per-name exposure is operable-MW share (operated view, EIA-860) - regulated fuel costs mostly pass through '
                      'fuel clauses; merchant names carry them in margin.')
    for key, fn, unit in (('henry_hub', 'MHHNGSP.xlsx', '$/MMBtu'), ('uranium', 'PURANUSDM.xlsx', '$/lb'),
                          ('transformer_ppi', 'PCU335311335311.xlsx', 'index Jun-1981=100')):
        ser, title = fred(fn)
        doc[key] = summarise(ser, unit)
        doc[key]['title'] = title
        if key == 'transformer_ppi':
            by = dict(ser)
            base = by.get('2020-01')
            doc[key]['change_since_2020_01_pct'] = round(100 * (ser[-1][1] / base - 1), 1) if base else None

    cols, dl, _ = eia_csv('Coal_shipments_to_the_electric_power_sector__price_by_plant_state.csv')
    state_q = {}
    for name, rec in dl.items():
        st = name.split(' : ')[-1].strip()
        if st == 'United States':
            state_q['US'] = rec['values']
        elif st in STATE_ABBR:
            state_q[STATE_ABBR[st]] = rec['values']
    qcols = [c for c in cols if any(c in v for v in state_q.values())]
    last_q = qcols[-1]
    prev_q = 'Q%s %d' % (last_q[1], int(last_q[-4:]) - 1)
    doc['coal_delivered'] = {'unit': '$/short ton', 'latest_quarter': last_q,
                             'us': {'latest': state_q['US'].get(last_q), 'year_ago': state_q['US'].get(prev_q),
                                    'quarterly': [[q, state_q['US'][q]] for q in qcols if q in state_q['US']]},
                             'by_state': {s: {'latest': v.get(last_q), 'year_ago': v.get(prev_q),
                                              'last_reported': max((q for q in v), key=lambda q: (q[-4:], q[1]))}
                                          for s, v in sorted(state_q.items()) if s != 'US'}}
    ycols, mk, _ = eia_csv('Market_average_price.csv')
    doc['coal_mine_price'] = {'unit': '$/short ton', 'years': ycols,
                              'series': {n: rec['values'] for n, rec in mk.items() if 'open market' in n or 'captive' in n}}

    # per-name exposure off the fleet layer
    fl = json.load(open(os.path.join(DATA, 'fleet.json'), encoding='utf-8')).get('tickers', {})
    names = {}
    for t, v in sorted(fl.items()):
        tot = v.get('operable_mw') or 0
        if not tot:
            continue
        bt = v.get('by_tech') or {}
        gas = sum(mw for k, mw in bt.items() if k.startswith('Gas'))
        coal = bt.get('Coal', 0.0)
        nuc = bt.get('Nuclear', 0.0)
        cs = collections.Counter()
        for u in (v.get('coal') or {}).get('units') or []:
            cs[u.get('state')] += u.get('mw') or 0
        w_now = [(state_q.get(s, {}).get(last_q), mw) for s, mw in cs.items()]
        w_prev = [(state_q.get(s, {}).get(prev_q), mw) for s, mw in cs.items()]
        wavg = lambda L: round(sum(p * m for p, m in L if p is not None) / sum(m for p, m in L if p is not None), 2) \
            if any(p is not None for p, _ in L) else None
        names[t] = {'operable_mw': tot, 'gas_share_pct': round(100 * gas / tot, 1), 'coal_share_pct': round(100 * coal / tot, 1),
                    'nuclear_share_pct': round(100 * nuc / tot, 1),
                    'coal_states_mw': {s: round(m) for s, m in cs.most_common()},
                    'coal_delivered_wavg': wavg(w_now), 'coal_delivered_wavg_year_ago': wavg(w_prev)}
    doc['names'] = names
    p = os.path.join(DATA, 'commodities.json')
    json.dump(doc, open(p, 'w', encoding='utf-8'), indent=1)
    hh, u, tp = doc['henry_hub'], doc['uranium'], doc['transformer_ppi']
    print(f"Henry Hub {hh['latest']} yoy {hh['yoy_pct']}% | uranium {u['latest']} yoy {u['yoy_pct']}% | transformer PPI {tp['latest']} "
          f"+{tp['change_since_2020_01_pct']}% since Jan-2020 | coal delivered US {doc['coal_delivered']['us']['latest']} ({last_q})")
    for t, v in names.items():
        if v['coal_share_pct'] or v['gas_share_pct'] > 30:
            print(f"  {t:5} gas {v['gas_share_pct']:5}%  coal {v['coal_share_pct']:5}%  coal $/t {v['coal_delivered_wavg']} (yr ago {v['coal_delivered_wavg_year_ago']})")
    print('wrote', p)


if __name__ == '__main__':
    main()
