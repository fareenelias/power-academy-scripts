# -*- coding: utf-8 -*-
r"""build_roe_drivers.py - ROE bridge Tier 3: WHY a utility under-earns its allowed ROE.
Actual vs test-year on rate base, O&M, sales and equity thickness, per rate case and per company.

  python E:\PowerAcademy\scripts\build_roe_drivers.py [--reports DIR] [--out data\roe_drivers.json]

Inputs (already on disk):
  data\reports\*_Report_<date>.xlsx   CapIQ workbooks:
        Past Rate Cases                    per opco x state x service: test-year end, authorized ROE,
                                           equity ratio, rate base, decision date
        Electric O&M Expenses / Sales Detail / Plant In Service  (and the Natural Gas versions)
                                           FERC Form 1/2 data summed to the holdco, last five years
  data\sec_opco.json                  10-K equity, LT debt, net utility plant per SEC registrant opco
  data\ferc_opco.json                 FERC Form 1 earned ROE per opco (Tier 2)

The four drivers, per latest decided case (decided 2015+) - EXPOSURES, not a reconciliation:
  1 Rate base lag   holdco net utility plant growth from the test-year end to the latest FY (FERC sheets;
                    opco-level from the 10-K when the opco is an SEC registrant and the test year is covered).
                    Gross ROE drag if none of it were in rates:  allowed ROE x (1 - NP_testyear / NP_latest).
                    Riders, trackers, formula rates, forward test years and CWIP-in-rates recover much of this,
                    so the drag is an UPPER bound - the case's own rider mechanisms decide how much is real.
  2 O&M creep       controllable electric O&M (total O&M less fuel, purchased & other power supply, and
                    transmission OPERATIONS, which are mostly RTO charges that pass through) - CAGR from
                    the test-year to the latest FY, against the growth in retail MWh and customers.
  3 Sales           retail MWh change since the test year (volumetric rates recover fixed cost on test-year
                    sales; decoupled jurisdictions are insulated - see the RRA decoupling flag).
  4 Equity thickness actual common equity / (equity + LT debt) from the opco 10-K vs the authorized equity
                    ratio. ROE drag if thicker:  allowed ROE x (1 - authorized / actual). Basis gap: the
                    authorized ratio often counts short-term debt / other capital, the 10-K ratio here does not.
"""
import os, re, sys, json, glob, datetime as dt
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TICKER_MAP = {'SJW': 'HTO', 'NEP': 'XIFR'}


def ticker_of(fn):
    m = re.search(r'(?:NYSE|NASDAQGS|NASDAQGM|NASDAQ|TSX|OTC)([A-Z0-9]+)_Report', os.path.basename(fn), re.I)
    return TICKER_MAP.get(m.group(1), m.group(1)) if m else None


def date_of(fn):
    m = re.search(r'_Report_(\d\d)-(\d\d)-(\d{4})', fn)
    return f'{m.group(3)}-{m.group(1)}-{m.group(2)}' if m else None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def series_sheet(ws):
    """label -> {year: value} for the CapIQ 'Last Five Years' layout (row 'Date Ended' holds the years)."""
    years, out = None, {}
    for r in ws.iter_rows(values_only=True):
        if not r or r[0] is None: continue
        lab = str(r[0]).strip()
        if lab == 'Date Ended':
            years = [(i, x.year) for i, x in enumerate(r) if hasattr(x, 'year')]
            continue
        if years:
            vals = {y: num(r[i]) for i, y in years if i < len(r)}
            if any(v is not None for v in vals.values()):
                out.setdefault(lab, vals)
    return out


def g(S, lab):
    return S.get(lab, {})


def sub(a, *bs):
    ys = set(a)
    out = {}
    for y in ys:
        v = a.get(y)
        if v is None: out[y] = None; continue
        for b in bs:
            v -= (b.get(y) or 0)
        out[y] = v
    return out


def electric(wb):
    if 'Electric O&M Expenses' not in wb.sheetnames:
        return None
    O = series_sheet(wb['Electric O&M Expenses'])
    fuel_labels = [k for k in O if re.search(r'Oper\s*:\s*Fuel', k)]
    fuel = {}
    for k in fuel_labels:
        for y, v in O[k].items():
            fuel[y] = (fuel.get(y) or 0) + (v or 0)
    tot = g(O, 'Total Electric Operation & Maintenance Expense')
    ctrl = sub(tot, fuel, g(O, 'Total Other Power Supply Expenses'), g(O, 'Transmission Total Operations Expense'))
    S = series_sheet(wb['Electric Sales Detail']) if 'Electric Sales Detail' in wb.sheetnames else {}
    P = series_sheet(wb['Electric Plant In Service - Sum']) if 'Electric Plant In Service - Sum' in wb.sheetnames else {}
    k = lambda d: {str(y): (None if v is None else round(v / 1000, 1)) for y, v in sorted(d.items())}   # $000 -> $M
    return {
        'om_total_m': k(tot), 'fuel_m': k(fuel), 'purchased_other_power_m': k(g(O, 'Total Other Power Supply Expenses')),
        'transmission_ops_m': k(g(O, 'Transmission Total Operations Expense')), 'om_controllable_m': k(ctrl),
        'om_distribution_m': k(g(O, 'Total Distribution O&M Expense')), 'om_ag_m': k(g(O, 'Total Adminstrative & General O&M Expense')),
        'retail_mwh_k': {str(y): (None if v is None else round(v / 1000)) for y, v in sorted(g(S, 'Total Retail Electric Volume, Total').items())},
        'retail_customers': {str(y): (None if v is None else int(v)) for y, v in sorted(g(S, 'Total Retail Electric Customers, Total').items())},
        'retail_revenue_m': k(g(S, 'Total Retail Electric Revenue')),
        'plant_in_service_m': k(g(P, 'Total Plant: Total')), 'net_utility_plant_m': k(g(P, 'Net Utility Plant: Total')),
        'cwip_m': k(g(P, 'Construction Work In Progress: Total')),
    }


def rate_cases(wb):
    if 'Past Rate Cases' not in wb.sheetnames:
        return []
    rows = list(wb['Past Rate Cases'].iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and r[0] == 'State'), None)
    if hi is None: return []
    top, subh = rows[hi], rows[hi + 1]
    # the two header rows: top names the block (requested / authorized), sub names the field
    cols, block = {}, ''
    for i in range(max(len(top), len(subh))):
        t = (top[i] if i < len(top) else None) or ''
        s = (subh[i] if i < len(subh) else None) or ''
        if 'Requested' in str(t): block = 'req'
        if 'Authorized' in str(t): block = 'auth'
        name = str(s or t).strip()
        cols[i] = (block, name)
    def col(block, rx):
        return next((i for i, (b, n) in cols.items() if b == block and re.search(rx, n, re.I)), None)
    ix = {'state': 0, 'company': 1, 'docket': 3, 'service': 4, 'case_type': 5,
          'auth_date': col('auth', r'^Date$'), 'decision': col('auth', r'Decision'),
          'auth_roe': col('auth', r'Return on Equity'), 'auth_eq': col('auth', r'Common Equity'),
          'auth_rb': col('auth', r'Rate Base'), 'ty_end': col('auth', r'Test Year'),
          'auth_inc': col('auth', r'Rate Increase'), 'req_roe': col('req', r'Return on Equity'),
          'req_inc': col('req', r'Rate Increase')}
    out = []
    for r in rows[hi + 2:]:
        if not r or not r[0] or not r[1]: continue
        gv = lambda k: r[ix[k]] if ix.get(k) is not None and ix[k] < len(r) else None
        d = gv('auth_date')
        ty = str(gv('ty_end') or '')
        m = re.match(r'(\d\d)/(\d{4})', ty)
        out.append({'state': r[0], 'company': r[1], 'docket': gv('docket'), 'service': gv('service'),
                    'case_type': gv('case_type'),
                    'decided': d.date().isoformat() if hasattr(d, 'date') else None, 'decision': gv('decision'),
                    'auth_roe': num(gv('auth_roe')), 'auth_equity_ratio': num(gv('auth_eq')),
                    'auth_rate_base_m': num(gv('auth_rb')), 'auth_increase_m': num(gv('auth_inc')),
                    'req_roe': num(gv('req_roe')), 'req_increase_m': num(gv('req_inc')),
                    'test_year_end': ty if m else None,
                    'test_year': (int(m.group(2)) if int(m.group(1)) >= 7 else int(m.group(2)) - 1) if m else None})
    return out


STATES = {'AL': 'Alabama', 'AR': 'Arkansas', 'AZ': 'Arizona', 'CA': 'California', 'CO': 'Colorado', 'CT': 'Connecticut',
          'FL': 'Florida', 'IL': 'Illinois', 'IN': 'Indiana', 'KS': 'Kansas', 'KY': 'Kentucky', 'LA': 'Louisiana',
          'MA': 'Massachusetts', 'MI': 'Michigan', 'MN': 'Minnesota', 'MO': 'Missouri', 'MS': 'Mississippi',
          'NH': 'New Hampshire', 'NM': 'New Mexico', 'NY': 'New York', 'OH': 'Ohio', 'OK': 'Oklahoma', 'OR': 'Oregon',
          'PA': 'Pennsylvania', 'SC': 'South Carolina', 'TX': 'Texas', 'VA': 'Virginia', 'WI': 'Wisconsin', 'WV': 'West Virginia'}
# jurisdictions whose AUTHORIZED capital structure routinely includes zero-cost capital (ADIT, customer deposits,
# ITC) - the authorized 'equity ratio' there is not comparable to a 10-K equity/(equity+LTD) ratio
ZERO_COST_CAP_STATES = {'Arkansas', 'Florida', 'Indiana', 'Michigan'}


def nname(s):
    s = re.sub(r'\b([A-Z]{2})\b', lambda m: STATES.get(m.group(1), m.group(1)), (s or '').strip())
    s = re.sub(r'\bSvc\b', 'Service', s)
    s = s.lower().replace('&', ' and ')
    s = re.sub(r'\bof\b', ' ', s)
    s = re.sub(r'\(.*?\)', ' ', s)
    s = re.sub(r'\b(company|co|inc|incorporated|llc|corporation|corp|the|dba|d/b/a)\b', ' ', s)
    s = re.sub(r'[^a-z]', '', s)
    return NAME_ALIAS.get(s, s)


# merged / renamed entities (same idea as FERC_ALIASES in the UI)
NAME_ALIAS = {'gulfpower': 'floridapowerandlight', 'westarenergy': 'evergykansascentral',
              'kansascitypowerandlight': 'evergymetro', 'westernmassachusettselectric': 'nstarelectric',
              'entergygulfstateslouisiana': 'entergylouisiana'}


def cagr(a, b, n):
    if a is None or b is None or a <= 0 or b <= 0 or n <= 0: return None
    return round(100 * ((b / a) ** (1 / n) - 1), 2)


def at(series, y):
    return series.get(str(y)) if series else None


def build(rdir, sec, ferc):
    cur = {}
    for f in glob.glob(os.path.join(rdir, '*_Report_*.xlsx')):
        t = ticker_of(f)
        if t and (t not in cur or date_of(f) > date_of(cur[t])): cur[t] = f
    out = {}
    for t in sorted(cur):
        wb = openpyxl.load_workbook(cur[t], read_only=True, data_only=True)
        E = electric(wb)
        cases = rate_cases(wb)
        wb.close()
        if not cases and not E:
            continue
        # latest decided case per company x state x service, decided 2015+
        latest = {}
        for c in cases:
            if not c['decided'] or c['decided'] < '2015-01-01' or c['auth_roe'] is None and c['auth_equity_ratio'] is None:
                continue
            k = (nname(c['company']), c['state'], c['service'])     # merged entities collapse onto the survivor
            if k not in latest or c['decided'] > latest[k]['decided']:
                latest[k] = c
        # SEC opcos (10-K) and FERC earned ROE, by normalized name
        sec_e = {nname(e['name']): e for e in ((sec.get(t) or {}).get('entities') or []) if e.get('role') == 'opco'}
        ferc_e = {nname(x.get('ferc_name')): x for x in (ferc.get(t) or [])}
        latest_fy = None
        if E:
            ys = [int(y) for y, v in E['net_utility_plant_m'].items() if v is not None]
            latest_fy = max(ys) if ys else None
        rows = []
        for c in sorted(latest.values(), key=lambda c: (c['company'], c['state'], c['service'])):
            r = dict(c)
            nm = nname(c['company'])
            se = sec_e.get(nm) or next((v for k, v in sec_e.items() if len(k) >= 8 and len(nm) >= 8 and (k in nm or nm in k)), None)
            fe = ferc_e.get(nm) or next((v for k, v in ferc_e.items() if len(k) >= 8 and len(nm) >= 8 and (k in nm or nm in k)), None)
            ty = c['test_year']
            drv = {}
            # 1 rate base lag (holdco FERC plant; opco 10-K when available)
            e_first = min((int(y) for y, v in E['net_utility_plant_m'].items() if v is not None), default=None) if E else None
            if E and ty and latest_fy and ty < latest_fy and (c['service'] or '').lower().startswith('electric'):
                y0 = max(ty, e_first)                      # test year older than the 5-year window -> lower bound
                np0, np1 = at(E['net_utility_plant_m'], y0), at(E['net_utility_plant_m'], latest_fy)
                if np0 and np1:
                    drv['rate_base'] = {'basis': 'holdco FERC net utility plant', 'from_year': y0, 'to_year': latest_fy,
                                        'lower_bound': y0 > ty,
                                        'growth_pct': round(100 * (np1 / np0 - 1), 1),
                                        'gross_roe_drag_bps': round(100 * c['auth_roe'] * (1 - np0 / np1)) if c['auth_roe'] else None}
            if se and ty:
                npo = se['metrics'].get('net_utility_plant') or {}
                y1 = max((int(y) for y, v in npo.items() if v is not None), default=None)
                if y1 and npo.get(str(ty)) and npo.get(str(y1)) and y1 > ty:
                    drv['rate_base_opco'] = {'basis': '10-K net utility plant (opco)', 'from_year': ty, 'to_year': y1,
                                             'growth_pct': round(100 * (npo[str(y1)] / npo[str(ty)] - 1), 1),
                                             'gross_roe_drag_bps': round(100 * c['auth_roe'] * (1 - npo[str(ty)] / npo[str(y1)])) if c['auth_roe'] else None}
            # 2 O&M creep vs 3 sales (holdco, electric cases)
            if E and ty and latest_fy and latest_fy > ty and (c['service'] or '').lower().startswith('electric'):
                ty0 = max(ty, e_first)
                n = latest_fy - ty0
                om0, om1 = at(E['om_controllable_m'], ty0), at(E['om_controllable_m'], latest_fy)
                s0, s1 = at(E['retail_mwh_k'], ty0), at(E['retail_mwh_k'], latest_fy)
                c0, c1 = at(E['retail_customers'], ty0), at(E['retail_customers'], latest_fy)
                if om0 and om1:
                    drv['om'] = {'basis': 'holdco controllable electric O&M (FERC)', 'from_year': ty0, 'to_year': latest_fy,
                                 'from_m': om0, 'to_m': om1, 'cagr_pct': cagr(om0, om1, n),
                                 'per_customer_cagr_pct': cagr(om0 / c0, om1 / c1, n) if c0 and c1 else None}
                if s0 and s1:
                    drv['sales'] = {'basis': 'holdco retail MWh (FERC)', 'from_year': ty0, 'to_year': latest_fy,
                                    'change_pct': round(100 * (s1 / s0 - 1), 1), 'cagr_pct': cagr(s0, s1, n),
                                    'customers_cagr_pct': cagr(c0, c1, n) if c0 and c1 else None}
            # 4 equity thickness (opco 10-K)
            if se and c['auth_equity_ratio']:
                m = se['metrics']
                y1 = max((int(y) for y, v in (m.get('equity') or {}).items() if v is not None), default=None)
                eq, ltd = (m.get('equity') or {}).get(str(y1)), (m.get('lt_debt') or {}).get(str(y1))
                if eq and ltd:
                    act = 100 * eq / (eq + ltd)
                    zc = c['state'] in ZERO_COST_CAP_STATES or c['auth_equity_ratio'] < 42
                    drv['equity'] = {'basis': '10-K equity / (equity + LT debt), opco', 'year': y1,
                                     'basis_warning': ('authorized structure likely includes zero-cost capital '
                                                       '(ADIT / customer deposits) - not comparable; excluded from the roll-up') if zc else None,
                                     'actual_ratio_pct': round(act, 1), 'authorized_ratio_pct': c['auth_equity_ratio'],
                                     'gap_pts': round(act - c['auth_equity_ratio'], 1),
                                     'roe_drag_bps': round(100 * c['auth_roe'] * (1 - c['auth_equity_ratio'] / act))
                                     if c['auth_roe'] and act > c['auth_equity_ratio'] else 0}
            # earned (FERC Form 1, Tier 2) for context
            if fe:
                ys = sorted((int(y) for y, v in fe['years'].items() if v.get('earned_roe_pct') is not None), reverse=True)
                if ys:
                    drv['earned'] = {'basis': 'FERC Form 1 earned ROE (total company)', 'year': ys[0],
                                     'earned_roe_pct': fe['years'][str(ys[0])]['earned_roe_pct'],
                                     'gap_bps': round(100 * (fe['years'][str(ys[0])]['earned_roe_pct'] - c['auth_roe'])) if c['auth_roe'] else None}
            r['years_since_test_year'] = (latest_fy or dt.date.today().year - 1) - ty if ty else None
            r['forward_test_year'] = bool(ty and latest_fy and ty > latest_fy - 0) and (c['decided'] or '')[:4] <= str(ty)
            r['sec_match'] = se['name'] if se else None
            r['ferc_match'] = fe.get('ferc_name') if fe else None
            r['drivers'] = drv
            rows.append(r)
        # company roll-up: authorized-rate-base-weighted averages of the exposures
        def wavg(key, fld):
            num_, den = 0.0, 0.0
            for r in rows:
                d = r['drivers'].get(key)
                w = r.get('auth_rate_base_m')
                if d and d.get(fld) is not None and w and not d.get('basis_warning'):
                    num_ += d[fld] * w; den += w
            return round(num_ / den) if den else None
        summ = {'cases': len(rows), 'latest_fy': latest_fy,
                'avg_years_since_test_year': round(sum(r['years_since_test_year'] for r in rows if r['years_since_test_year'] is not None) /
                                                   max(1, sum(1 for r in rows if r['years_since_test_year'] is not None)), 1) if rows else None,
                'rb_weighted_rate_base_drag_bps': wavg('rate_base', 'gross_roe_drag_bps'),
                'rb_weighted_equity_drag_bps': wavg('equity', 'roe_drag_bps'),
                'rb_weighted_earned_gap_bps': wavg('earned', 'gap_bps')}
        if E and latest_fy:
            y0 = min(int(y) for y, v in E['om_controllable_m'].items() if v is not None)
            n = latest_fy - y0
            summ['five_year'] = {'from_year': y0, 'to_year': latest_fy,
                                 'om_controllable_cagr_pct': cagr(at(E['om_controllable_m'], y0), at(E['om_controllable_m'], latest_fy), n),
                                 'retail_mwh_cagr_pct': cagr(at(E['retail_mwh_k'], y0), at(E['retail_mwh_k'], latest_fy), n),
                                 'customers_cagr_pct': cagr(at(E['retail_customers'], y0), at(E['retail_customers'], latest_fy), n),
                                 'net_plant_cagr_pct': cagr(at(E['net_utility_plant_m'], y0), at(E['net_utility_plant_m'], latest_fy), n),
                                 'retail_revenue_cagr_pct': cagr(at(E['retail_revenue_m'], y0), at(E['retail_revenue_m'], latest_fy), n)}
        flags = []
        fy_ = summ.get('five_year') or {}
        if fy_.get('om_controllable_cagr_pct') is not None and fy_.get('retail_mwh_cagr_pct') is not None and \
                fy_['om_controllable_cagr_pct'] - fy_['retail_mwh_cagr_pct'] > 3:
            flags.append(f"controllable O&M grew {fy_['om_controllable_cagr_pct']}%/yr vs retail MWh {fy_['retail_mwh_cagr_pct']}%/yr ({fy_['from_year']}-{fy_['to_year']})")
        if summ['rb_weighted_rate_base_drag_bps'] and summ['rb_weighted_rate_base_drag_bps'] >= 150:
            flags.append(f"rate base has outrun test years: gross drag ~{summ['rb_weighted_rate_base_drag_bps']} bps before riders")
        if summ['rb_weighted_equity_drag_bps'] and summ['rb_weighted_equity_drag_bps'] >= 50:
            flags.append(f"opco equity thicker than authorized: ~{summ['rb_weighted_equity_drag_bps']} bps drag")
        stale = [r for r in rows if r['years_since_test_year'] and r['years_since_test_year'] >= 4]
        if stale:
            flags.append(f"{len(stale)} of {len(rows)} jurisdictions on test years 4+ years old")
        out[t] = {'as_of': date_of(cur[t]), 'electric_series': E, 'cases': rows, 'summary': summ, 'flags': flags}
    return out


def main():
    a = sys.argv[1:]
    rdir = a[a.index('--reports') + 1] if '--reports' in a else os.path.join(BASE, 'data', 'reports')
    data = a[a.index('--data') + 1] if '--data' in a else os.path.join(BASE, 'data')
    outp = a[a.index('--out') + 1] if '--out' in a else os.path.join(data, 'roe_drivers.json')
    sec = json.load(open(os.path.join(data, 'sec_opco.json'), encoding='utf-8'))['tickers']
    ferc = json.load(open(os.path.join(data, 'ferc_opco.json'), encoding='utf-8'))['opcos']
    names = build(rdir, sec, ferc)
    doc = {'_schema_version': '1.0', '_generated': dt.date.today().isoformat(),
           '_method': __doc__.split('The four drivers')[1].split('"""')[0].strip(),
           '_units': '$M unless noted; MWh in thousands; ROE drags in basis points',
           'names': names}
    json.dump(doc, open(outp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'wrote {outp}: {len(names)} names')
    for t, v in names.items():
        s = v['summary']
        print(f"  {t:5} cases {s['cases']:3}  TY age {s['avg_years_since_test_year']}  RBdrag {s['rb_weighted_rate_base_drag_bps']}  "
              f"EQdrag {s['rb_weighted_equity_drag_bps']}  earned gap {s['rb_weighted_earned_gap_bps']} | {'; '.join(v['flags'])[:160]}")


if __name__ == '__main__':
    main()
