# -*- coding: utf-8 -*-
"""extract_ownership.py - CapIQ ownership / activism sheets -> data\\ownership.json (Activist tab).

Reads, per workbook in data\\reports (newest per ticker):
  'Investor Activism Summary'  activist holders (Data block), key developments, campaign history;
                               the sheet prints 'There are no campaigns' when the 5-year window is
                               empty; absent sheet = CapIQ did not offer it for the name (recorded)
  'Top Holders'                current top holders with CapIQ's 'Institution is an Activist' flag
  'Ownership Activity'         quarter flow: institutions / individuals & insiders / other strategic
  'Insider Activity'           aggregates (3m/1y/5y buy vs sell) + the transaction list (1 year)

Everything is as printed; nothing is inferred. The Activist tab renders this file and keeps
Fareen's free-text notes beside it.

  python extract_ownership.py [--dry]
"""
import io, os, re, sys, glob, json, collections, datetime

DATA_DIR = r'E:\PowerAcademy\data' if os.name == 'nt' else \
    os.path.join(os.path.expanduser('~'), 'mnt', 'PowerAcademy', 'data')
try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '--break-system-packages', '-q'])
    import openpyxl


def ticker_from_filename(f):
    b = re.sub(r'_Report_.*$', '', os.path.splitext(os.path.basename(f))[0])
    m = re.search(r'(NYSE|NASDAQGS|NASDAQGM|NASDAQCM|TSX)([A-Z]+)$', b)
    return m.group(2) if m else None


def wdate(f):
    m = re.search(r'_Report_(\d{2})-(\d{2})-(\d{4})', os.path.basename(f))
    return f'{m.group(3)}-{m.group(1)}-{m.group(2)}' if m else ''


def num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(',', '')
    if s.upper() in ('NA', 'N/A', '', '-', 'NM'): return None
    try: return float(s)
    except ValueError: return None


def dt(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d')
    s = str(v).strip(); return s[:10] if re.match(r'\d{4}-\d{2}-\d{2}', s) else (s or None)


def txt(v):
    if v is None: return None
    s = re.sub(r'\s+', ' ', str(v)).strip()
    return s or None


def rows_of(wb, name):
    return [tuple(r) for r in wb[name].iter_rows(values_only=True)] if name in wb.sheetnames else None


def block_after(rows, header_first_cell, stop_blank=True):
    """Rows following the row whose first cell == header_first_cell, until a blank first cell."""
    out, on = [], False
    for r in rows:
        c0 = txt(r[0]) if r else None
        if not on:
            if c0 == header_first_cell: on = True; hdr = [txt(x) for x in r]
            continue
        if c0 is None or c0 == '' : 
            if stop_blank: break
            continue
        out.append(r)
    return out


def parse_activism(rows):
    if rows is None: return None
    a = collections.OrderedDict()
    head = ' '.join(txt(r[0]) or '' for r in rows[:12] if r and r[0])
    a['period'] = next((txt(r[0]).split(':', 1)[1].strip() for r in rows[:8] if r and txt(r[0]) and txt(r[0]).startswith('Period')), None)
    if 'no campaign' in head.lower():
        a['no_campaigns'] = True; a['activist_holders'] = []; a['key_developments'] = []; a['campaigns'] = []
        return a
    a['no_campaigns'] = False
    hold = []
    for r in block_after(rows, 'Activist'):
        if txt(r[0]) == 'Total': break
        hold.append(collections.OrderedDict([('activist', txt(r[0])), ('shares', num(r[1])), ('pct_outstanding', num(r[2])),
                                             ('market_value_m', num(r[3])), ('position_date', dt(r[4]))]))
    a['activist_holders'] = hold
    a['key_developments'] = [collections.OrderedDict([('date', dt(r[0])), ('headline', txt(r[1]))])
                             for r in block_after(rows, 'Date') if dt(r[0])]
    camps = []
    for r in block_after(rows, 'Campaign ID'):
        camps.append(collections.OrderedDict([('campaign_id', txt(r[0])), ('launched', dt(r[1])), ('activists', txt(r[2])),
                                              ('status', txt(r[3])), ('ended', dt(r[4])), ('tactics', txt(r[5])), ('objectives', txt(r[6]))]))
    a['campaigns'] = camps
    return a


def parse_top_holders(rows):
    if rows is None: return None
    out = []
    for r in block_after(rows, 'Holder'):
        if not txt(r[0]) or txt(r[0]).lower().startswith('total'): continue
        out.append(collections.OrderedDict([('holder', txt(r[0])), ('shares', num(r[2])), ('pct_outstanding', num(r[3])),
                                            ('market_value_m', num(r[4])), ('position_date', dt(r[5])),
                                            ('is_activist', (txt(r[6]) or '').lower() == 'yes')]))
    return out


def parse_ownership_activity(rows):
    if rows is None: return None
    o = collections.OrderedDict()
    period = None
    for r in rows:
        if r and txt(r[0]) and re.match(r'[A-Z][a-z]{2}-\d{2}-\d{4} To', txt(r[0])):
            period = txt(r[0]); break
    o['period'] = period
    def grp(rows, col, key):
        d = collections.OrderedDict()
        for r in rows:
            c = txt(r[col]) if len(r) > col else None
            if c in ('Total Positions', 'New Positions', 'Increased Positions', 'Decreased Positions', 'Sold Out Positions'):
                d[c.lower().replace(' ', '_')] = dict(holders=num(r[col + 1]) if len(r) > col + 1 else None, shares=num(r[col + 2]) if len(r) > col + 2 else None)
        return d
    # institutions in cols 0-2, individuals/insiders in cols 3-5, then the 'Other Strategic' block re-uses cols 0-2
    strat_start = next((i for i, r in enumerate(rows) if r and txt(r[0]) and txt(r[0]).startswith('Other Strategic')), None)
    top = rows[:strat_start] if strat_start is not None else rows
    o['institutions'] = grp(top, 0, 'inst')
    o['individuals_insiders'] = grp(top, 3, 'ind')
    o['other_strategic'] = grp(rows[strat_start:], 0, 'strat') if strat_start is not None else {}
    return o


def parse_insider(rows):
    if rows is None: return None
    ins = collections.OrderedDict()
    agg = collections.OrderedDict()
    on = False
    for r in rows:
        c0 = txt(r[0])
        if c0 == 'Aggregates': on = True; continue
        if on:
            if c0 == 'Last 3 Months' or (c0 and c0.startswith('Aggregates are')): 
                if c0.startswith('Aggregates are'): break
                continue
            if c0 and len(r) > 3: agg[c0] = dict(last_3m=txt(r[1]), last_1y=txt(r[2]), last_5y=txt(r[3]))
    ins['aggregates'] = agg
    tx = []
    holder = None
    hdr = next((r for r in rows if r and txt(r[0]) == 'Holder Name'), None)
    if hdr:
        H = {txt(h): i for i, h in enumerate(hdr) if txt(h)}
        def col(r, name):
            i = H.get(name)
            if i is None: i = next((v for k, v in H.items() if k.startswith(name)), None)
            return r[i] if i is not None and i < len(r) else None
        for r in block_after(rows, 'Holder Name', stop_blank=False):
            if not dt(col(r, 'Trade Date Range')): continue
            h = txt(r[0])
            if h: holder = h
            tx.append(collections.OrderedDict([('holder', holder), ('trade_date', dt(col(r, 'Trade Date Range'))),
                ('security', txt(col(r, 'Security Type'))), ('shares', num(col(r, 'Transacted Shares'))),
                ('value', num(col(r, 'Transaction Value Range'))), ('type', txt(col(r, 'Transaction Type'))),
                ('sec_code', txt(col(r, 'SEC Transaction Code'))), ('price', txt(col(r, 'Price Range ($)') or col(r, 'Price Range (C$)') or col(r, 'Price Range'))),
                ('end_shares', num(col(r, 'End of Filing Shares'))), ('pct_change', num(col(r, '% Change'))),
                ('filed', dt(col(r, 'Filed Date'))), ('source', txt(col(r, 'Source')))]))
    ins['transactions'] = tx
    ins['n_transactions'] = len(tx)
    return ins


def main():
    dry = '--dry' in sys.argv
    files = sorted(glob.glob(os.path.join(DATA_DIR, 'reports', '*.xlsx')))
    best = {}
    for f in files:
        t = ticker_from_filename(f)
        if t and (t not in best or wdate(f) > wdate(best[t])): best[t] = f
    out = collections.OrderedDict()
    out['_schema_version'] = '1.0'
    out['_generated'] = datetime.date.today().isoformat()
    out['_source'] = "CapIQ workbook sheets: Investor Activism Summary / Top Holders / Ownership Activity / Insider Activity - as printed"
    out['_note'] = ("activism = null means CapIQ offered no activism sheet for the name (AWR, WTRG, GWRS, MSEX, POR, TLN, YORW, VST, XIFR on "
                    "the 09-23 pull) - not 'no activism'. no_campaigns = true is CapIQ's own statement for the 5-year window.")
    out['tickers'] = collections.OrderedDict()
    for t in sorted(best):
        f = best[t]
        wb = openpyxl.load_workbook(f, read_only=True)
        rec = collections.OrderedDict([('source_workbook', os.path.basename(f)), ('workbook_date', wdate(f))])
        rec['activism'] = parse_activism(rows_of(wb, 'Investor Activism Summary'))
        rec['top_holders'] = parse_top_holders(rows_of(wb, 'Top Holders'))
        rec['ownership_activity'] = parse_ownership_activity(rows_of(wb, 'Ownership Activity'))
        rec['insider'] = parse_insider(rows_of(wb, 'Insider Activity'))
        wb.close()
        out['tickers'][t] = rec
        a = rec['activism']; th = rec['top_holders'] or []
        astr = 'n/a' if a is None else ('none' if a['no_campaigns'] else '%d campaigns/%d holders' % (len(a['campaigns']), len(a['activist_holders'])))
        print('%-5s activism=%-22s top_holders=%2d activist_flag=%s  insider_tx=%s  own_period=%s' % (
            t, astr, len(th), [h['holder'] for h in th if h['is_activist']], (rec['insider'] or {}).get('n_transactions'),
            (rec['ownership_activity'] or {}).get('period')))
    if dry: return
    p = os.path.join(DATA_DIR, 'ownership.json')
    io.open(p, 'w', encoding='utf-8').write(json.dumps(out, indent=1, ensure_ascii=False)); json.load(io.open(p, encoding='utf-8'))
    print('wrote', p)


if __name__ == '__main__':
    main()
