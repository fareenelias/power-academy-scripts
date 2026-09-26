# -*- coding: utf-8 -*-
r"""build_succession.py - leadership changes and succession-risk flags per name, from the CapIQ workbooks.

  python E:\PowerAcademy\scripts\build_succession.py [--reports DIR] [--out data\succession.json]

Inputs (all CapIQ, already on disk):
  data\reports\*_Report_<date>.xlsx            People (Top Executives / Board) + Corporate Governance
  data\reports\_archive\*_Report_<date>.xlsx   older pulls of the same workbooks -> what CHANGED between pulls

Why not the People 'Status' column: every pull is exported with View = Current, so Status reads 'Current'
for every row (615 of 615 on 2026-09-23). Departures are found by DIFFING the current pull against the
oldest archived pull instead. Re-export People with View = 'All' to get CapIQ's own former-officer rows.

Flags (heuristics, stated as such on screen; thresholds live here, not in the UI):
  ceo_age_63_plus       CEO age >= 63 (typical utility CEO retirement window 63-65)
  ceo_tenure_10y_plus   CEO in role >= 10 years (bio-parsed start year)
  ceo_new_lt_2y         CEO in role < 2 years (transition still settling)
  cfo_new_lt_1y         CFO in role < 1 year
  no_named_heir         no President / COO title held by someone other than the CEO
  combined_chair_ceo    CEO also chairs the board
  csuite_departure      a CEO / CFO / COO / President / GC / CAO title-holder in the older pull is gone now
  csuite_arrival        a C-suite title appears that was not in the older pull
  csuite_title_change   a C-suite title-holder's title changed between pulls (e.g. CFO moved to another role)
  opco_leader_departure a 'President / CEO of <opco or business>' is no longer listed
  directors_at_retirement  directors within 1 year of the board's mandatory retirement age
"""
import os, re, sys, json, glob, datetime as dt
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TICKER_MAP = {'SJW': 'HTO', 'NEP': 'XIFR'}
TODAY = dt.date.today()
CSUITE = re.compile(r'\b(CEO|Chief Executive|CFO|Chief Financial|COO|Chief Operating|President|General Counsel|'
                    r'Chief Legal|Chief Accounting|Controller|Chief Administrative)\b', re.I)
MONTHS = 'January|February|March|April|May|June|July|August|September|October|November|December'


def ticker_of(fn):
    m = re.search(r'(?:NYSE|NASDAQGS|NASDAQGM|NASDAQ|TSX|OTC)([A-Z0-9]+)_Report', os.path.basename(fn), re.I)
    return TICKER_MAP.get(m.group(1), m.group(1)) if m else None


def date_of(fn):
    m = re.search(r'_Report_(\d\d)-(\d\d)-(\d{4})', fn)
    return f'{m.group(3)}-{m.group(1)}-{m.group(2)}' if m else None


def clean(v):
    if v is None: return None
    s = str(v).strip()
    return None if s.upper() in ('', 'NA', 'N/A', 'NONE') else s


def to_int(v):
    try:
        n = int(float(str(v).strip())); return n if n > 0 else None
    except (TypeError, ValueError):
        return None


def people(wb):
    """(execs, board) with columns located from each section's header row, not fixed offsets."""
    if 'People' not in wb.sheetnames:
        return None, None
    execs, board, mode, hdr = [], [], None, None
    for r in wb['People'].iter_rows(values_only=True):
        c0 = clean(r[0]) if r else None
        if c0 == 'Top Executives': mode, hdr = 'exec', None; continue
        if c0 == 'Board of Directors': mode, hdr = 'board', None; continue
        if c0 in ('Other Board Members', 'Professionals', 'Takeover Defenses') or re.search(r'News|Latest Holders|exports are limited|Contact information', c0 or ''): mode = None; continue
        if c0 == 'Average' or not mode or not c0: continue
        if c0 == 'Name':
            hdr = {clean(h): i for i, h in enumerate(r) if clean(h)}; continue
        if not hdr or re.search(r'\d', c0) or len(c0.split()) < 2 or len(c0) > 60: continue
        g = lambda k: r[hdr[k]] if k in hdr and hdr[k] < len(r) else None
        rec = {'name': c0, 'title': clean(g('Role')), 'status': clean(g('Status')), 'age': to_int(g('Age')),
               'bio': clean(g('Biography'))}
        if mode == 'board':
            rec.update({'begin_year': to_int(g('Begin Year')), 'tenure': clean(g('Tenure'))})
            board.append(rec)
        else:
            execs.append(rec)
    return execs, board


def governance(wb):
    out = {}
    if 'Corporate Governance' not in wb.sheetnames:
        return out
    want = {'Retirement Age for Directors (years)': 'director_retirement_age',
            'CEO Succession Plan?': 'ceo_succession_plan_disclosed', 'Classified Board?': 'classified_board'}
    for r in wb['Corporate Governance'].iter_rows(values_only=True):
        for i, c in enumerate(r):
            if c in want:
                v = next((x for x in r[i + 1:] if x is not None and str(x).strip() != ''), None)
                out[want[c]] = to_int(v) if want[c] == 'director_retirement_age' else v
    return out


def norm(name):
    s = re.sub(r'"[^"]*"', '', name or '')
    s = re.sub(r'\b(Jr|Sr|II|III|IV|Dr|Mr|Ms|Mrs)\b\.?', '', s)
    parts = re.sub(r'[^A-Za-z ]', ' ', s).lower().split()
    return (parts[0] + ' ' + parts[-1]) if len(parts) >= 2 else ' '.join(parts)


def role_start(bio, role_rx, company=''):
    """Year the bio says this person STARTED in the role. The bio is cut into clauses (', and has been',
    '; ', ' and serves' ...) so a date binds to the title in ITS clause - 'President since 2012 and has been
    its Chief Executive Officer since 2013' gives 2013. Clauses with 'until' / 'had been' / 'was' / 'served'
    (past roles) are skipped. A clause naming the company wins over one that names another company; among
    equals the LATEST start is taken (a promotion into the top job is the last 'since').
    Returns (year, clause) or (None, None)."""
    if not bio:
        return None, None
    key = [w for w in re.sub(r'[^A-Za-z ]', ' ', company).split()
           if w.lower() not in ('the', 'inc', 'corp', 'corporation', 'company', 'co', 'group', 'lp', 'energy',
                                'american', 'utilities', 'and', 'holdings')][:1]
    key = key[0].lower() if key else None
    cands = []
    for sent in re.split(r'(?<=[.])\s+', bio):
        for cl in re.split(r';|,?\s+and\s+(?=(?:has|have|serves|served|is|was|also|became|had)\b)|,\s+(?=(?:has|serves|served|is|was|also|became)\b)', sent):
            if not re.search(role_rx, cl, re.I): continue
            if re.search(r'\b(until|had been|was|served|former|previously|prior to|retired)\b', cl, re.I): continue
            if re.search(r'\b(Vice|Deputy|Assistant|Interim|Acting)\b[^.]{0,20}(Chief|CEO|CFO)', cl): continue
            rp = re.search(role_rx, cl, re.I).start()          # the date must FOLLOW the title
            m = re.search(rf'\b(?:since|effective|as of)\s+(?:(?:{MONTHS})\s+)?(?:\d{{1,2}},?\s*)?((?:19|20)\d\d)', cl[rp:])
            if not m: continue
            y = int(m.group(1))
            if y > TODAY.year: continue
            names_other = re.search(r'\b(?:of|at)\s+([A-Z][A-Za-z&.]+(?:\s+[A-Z][A-Za-z&.]+)*)', cl)
            mine = bool(key and key in cl.lower())
            other = bool(names_other and not mine and not re.match(r'(the Board|Directors|the Company)', names_other.group(1)))
            if other: continue
            cands.append((mine, y, cl.strip()[:300]))
    if not cands:
        return None, None
    cands.sort(key=lambda c: (c[0], c[1]))
    return cands[-1][1], cands[-1][2]


def _tnorm(t):
    t = (t or '').lower().replace('&', ' and ')
    t = re.sub(r'\bexecutive vice president\b', 'evp', t); t = re.sub(r'\bexecutive vp\b', 'evp', t)
    t = re.sub(r'\bsenior vice president\b', 'svp', t); t = re.sub(r'\bsenior vp\b', 'svp', t)
    return re.sub(r'[^a-z]|\bof\b|\band\b', '', re.sub(r'\b(of|and|the)\b', ' ', t))


def pick(execs, rx, exclude=None):
    for p in execs or []:
        if p['title'] and re.search(rx, p['title'], re.I) and p['name'] != exclude:
            return p
    return None


def build(reports_dir):
    cur = {}
    for f in glob.glob(os.path.join(reports_dir, '*_Report_*.xlsx')):
        t = ticker_of(f)
        if t and (t not in cur or date_of(f) > date_of(cur[t])): cur[t] = f
    arch = {}
    for f in glob.glob(os.path.join(reports_dir, '_archive', '*_Report_*.xlsx')):
        t = ticker_of(f)
        if t and (t not in arch or date_of(f) < date_of(arch[t])): arch[t] = f   # OLDEST pull = widest diff
    out = {}
    for t in sorted(cur):
        wb = openpyxl.load_workbook(cur[t], read_only=True, data_only=True)
        execs, board = people(wb)
        gov = governance(wb)
        # company name: file name 'AmericanElectricPowerCompany,Inc.NASDAQGSAEP_Report...' -> words
        company = re.sub(r'([a-z])([A-Z])', r'\1 \2', re.split(r'(?:NYSE|NASDAQ|TSX)', os.path.basename(cur[t]))[0])
        if False:
            for r in wb['People'].iter_rows(min_row=1, max_row=3, values_only=True):
                if r and r[0] and '|' in str(r[0]): company = str(r[0]).split('|')[0].strip(); break
        wb.close()
        people_src = date_of(cur[t])
        old = None
        if t in arch:
            wb = openpyxl.load_workbook(arch[t], read_only=True, data_only=True)
            old = people(wb); wb.close()
        if execs is None and old and old[0] is not None:        # current pull lacks People (HTO) -> older pull
            execs, board = old
            people_src = date_of(arch[t]) + ' (current pull has no People sheet)'
            old = None
        execs = execs or []; board = board or []
        ceo = pick(execs, r'\bCEO\b|Chief Executive')
        cfo = pick(execs, r'\bCFO\b|Chief Financial')
        heir = next((p for p in execs if p is not ceo and p['title'] and
                     re.search(r'^(?!.*\bof\b).*\b(President|COO|Chief Operating)\b', p['title'], re.I)), None)
        rec = {'as_of': date_of(cur[t]), 'people_source_date': people_src, 'governance': gov,
               'ceo': None, 'cfo': None, 'heir_candidates': [], 'changes': None, 'flags': [],
               'directors_near_retirement': []}
        for key, p, rx in (('ceo', ceo, r'\bCEO\b|Chief Executive'), ('cfo', cfo, r'\bCFO\b|Chief Financial')):
            if not p: continue
            y, ev = role_start(p['bio'], rx, company)
            rec[key] = {'name': p['name'], 'title': p['title'], 'age': p['age'], 'role_start_year': y,
                        'years_in_role': (TODAY.year - y) if y else None, 'evidence': ev}
        rec['heir_candidates'] = [{'name': p['name'], 'title': p['title'], 'age': p['age']}
                                  for p in execs if p is not ceo and p['title'] and
                                  re.search(r'\b(President|COO|Chief Operating)\b', p['title'], re.I) and
                                  not re.search(r'\bof\b|\bVice\b', p['title'], re.I)][:4]
        opco_presidents = [{'name': p['name'], 'title': p['title'], 'age': p['age']} for p in execs
                           if p['title'] and re.search(r'President.*\bof\b', p['title'], re.I) and not re.search(r'Vice', p['title'], re.I)]
        rec['opco_presidents'] = opco_presidents[:12]
        # ---- changes between pulls
        if old and old[0] is not None:
            o_ex = {norm(p['name']): p for p in old[0]}
            n_ex = {norm(p['name']): p for p in execs}
            o_bd = {norm(p['name']): p for p in (old[1] or [])}
            n_bd = {norm(p['name']): p for p in board}
            rec['changes'] = {
                'compared_to': date_of(arch[t]),
                'exec_departed': [{'name': o_ex[k]['name'], 'title': o_ex[k]['title']} for k in o_ex if k not in n_ex],
                'exec_arrived': [{'name': n_ex[k]['name'], 'title': n_ex[k]['title']} for k in n_ex if k not in o_ex],
                'title_changed': [{'name': n_ex[k]['name'], 'from': o_ex[k]['title'], 'to': n_ex[k]['title']}
                                  for k in n_ex if k in o_ex and (o_ex[k]['title'] or '') != (n_ex[k]['title'] or '')],
                'board_departed': [o_bd[k]['name'] for k in o_bd if k not in n_bd],
                'board_joined': [n_bd[k]['name'] for k in n_bd if k not in o_bd],
            }
        # ---- flags
        F = rec['flags']
        c = rec['ceo']
        if c and c['age'] and c['age'] >= 63:
            F.append({'id': 'ceo_age_63_plus', 'text': f"CEO {c['name']} is {c['age']}"})
        if c and c['years_in_role'] is not None and c['years_in_role'] >= 10:
            F.append({'id': 'ceo_tenure_10y_plus', 'text': f"CEO in role since {c['role_start_year']} (bio)"})
        if c and c['years_in_role'] is not None and c['years_in_role'] < 2:
            F.append({'id': 'ceo_new_lt_2y', 'text': f"CEO in role since {c['role_start_year']} (bio)"})
        f_ = rec['cfo']
        if f_ and f_['years_in_role'] is not None and f_['years_in_role'] < 1:
            F.append({'id': 'cfo_new_lt_1y', 'text': f"CFO {f_['name']} in role since {f_['role_start_year']} (bio)"})
        if c and not rec['heir_candidates']:
            F.append({'id': 'no_named_heir', 'text': 'no President / COO title held by someone other than the CEO'})
        if c and re.search(r'Chair', c['title'] or '', re.I):
            F.append({'id': 'combined_chair_ceo', 'text': 'CEO also chairs the board'})
        if rec['changes']:
            dep = [x for x in rec['changes']['exec_departed'] if CSUITE.search(x['title'] or '') and not re.search(r'\bof\b', x['title'] or '')]
            arr = [x for x in rec['changes']['exec_arrived'] if CSUITE.search(x['title'] or '') and not re.search(r'\bof\b', x['title'] or '')]
            if dep:
                F.append({'id': 'csuite_departure', 'text': 'no longer listed since ' + rec['changes']['compared_to'] + ': ' +
                          '; '.join(f"{x['name']} ({x['title']})" for x in dep)})
            moved = [x for x in rec['changes']['title_changed']
                     if any(CSUITE.search(v or '') and not re.search(r'\bof\b', v or '') for v in (x['from'], x['to']))
                     and _tnorm(x['from']) != _tnorm(x['to'])]
            if moved:
                F.append({'id': 'csuite_title_change', 'text': 'title changed since ' + rec['changes']['compared_to'] + ': ' +
                          '; '.join(f"{x['name']} ({x['from']} -> {x['to']})" for x in moved)})
            opl = [x for x in rec['changes']['exec_departed']
                   if re.search(r'\b(President|CEO|Chief Executive)\b', x['title'] or '') and re.search(r'\bof\b', x['title'] or '')
                   and not re.search(r'Vice', x['title'] or '')]
            if opl:
                F.append({'id': 'opco_leader_departure', 'text': 'opco / business-unit head no longer listed since ' +
                          rec['changes']['compared_to'] + ': ' + '; '.join(f"{x['name']} ({x['title']})" for x in opl)})
            if arr:
                F.append({'id': 'csuite_arrival', 'text': 'newly listed since ' + rec['changes']['compared_to'] + ': ' +
                          '; '.join(f"{x['name']} ({x['title']})" for x in arr)})
        ra = gov.get('director_retirement_age')
        if ra:
            near = [{'name': p['name'], 'age': p['age']} for p in board if p['age'] and p['age'] >= ra - 1]
            rec['directors_near_retirement'] = near
            if near:
                F.append({'id': 'directors_at_retirement', 'text': f"{len(near)} director(s) within a year of the age-{ra} retirement policy"})
        rec['counts'] = {'executives': len(execs), 'directors': len(board)}
        out[t] = rec
    return out


def main():
    a = sys.argv[1:]
    rdir = a[a.index('--reports') + 1] if '--reports' in a else os.path.join(BASE, 'data', 'reports')
    outp = a[a.index('--out') + 1] if '--out' in a else os.path.join(BASE, 'data', 'succession.json')
    names = build(rdir)
    doc = {'_schema_version': '1.0', '_generated': TODAY.isoformat(),
           '_source': 'S&P Capital IQ company workbooks (People, Corporate Governance), current pull vs oldest archived pull',
           '_method': __doc__.split('Flags (heuristics')[1].split('"""')[0].strip() if 'Flags (heuristics' in __doc__ else '',
           '_status_note': 'People sheet exported with View = Current, so the Status column is uniformly Current; '
                           'departures are found by diffing pulls. Role start years are parsed from CapIQ biographies '
                           '(evidence sentence kept) - check the sentence before relying on it.',
           'names': names}
    json.dump(doc, open(outp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    nf = sum(len(v['flags']) for v in names.values())
    print(f'wrote {outp}: {len(names)} names, {nf} flags')
    for t, v in names.items():
        c = v['ceo'] or {}
        print(f"  {t:5} CEO {str(c.get('name'))[:28]:28} age {c.get('age')} since {c.get('role_start_year')}  "
              f"flags: {', '.join(f['id'] for f in v['flags'])}")


if __name__ == '__main__':
    main()
