# -*- coding: utf-8 -*-
r"""apply_precedents_review.py - apply Fareen's review of PowerAcademy_review_2026-09-25_precedents(-FE comments).xlsx
to data\precedents.json (backup precedents.json.bak-preFEreview-20260925).

  python scripts\apply_precedents_review.py --review "E:\PowerAcademy\PowerAcademy_review_2026-09-25_precedents-FE comments.xlsx"

  1_Weak_links / 2_Verify_flags : Confirm -> _verified[field]='human', confidence 'verified', verify ok
                                  Replace -> new URL, verified; Remove -> link cleared, _needs_find
  3_Stale_flags                 : confidence entries on empty fields dropped
  4_Proxy_vs_file ('Proxy value'): the proxy's figures replace the deal facts (EDITS below, each with its
                                  basis), every extracted deal gets a `proxy` block (filing, fees, advisors
                                  as named in the proxy), and advisors the proxy names are added.
"""
import os, sys, json, shutil, re
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIELD = {'Press release (acquirer)': 'press_release_acquirer', 'Press release (target)': 'press_release_target',
         'Press release': 'press_release', 'Investor deck': 'deck', 'Fairness opinion': 'fairness_opinion',
         'Merger agreement': 'agreement', '8-K / filing': 'filing', 'transcript': 'transcript'}
# sheet-1 rows that are paste errors (URL belongs to another deal) -> skipped, sheet-2 value used instead
SKIP = {('1_Weak_links', 'cpl_floridaprogress_1999', 'deck')}
TODAY = '2026-09-25'

# 4_Proxy_vs_file decisions ('Proxy value' on every row) -> concrete field edits. Source = the merger proxy.
EDITS = {
 'blackhills_northwestern_2025': {'offer_px': 60.39, 'premium_pct': 9.0, 'premium_basis': "1-day, NWE $55.39 close 2025-08-15 (0.98 x BKH $61.62) - Goldman opinion"},
 'dominion_scana_2018': {'offer_px': 54.23, 'premium_pct': 36.3, 'premium_basis': '1-day, SCANA 12/29/2017 close (proxy); 29.7% on 30-day VWAP'},
 'firstenergy_allegheny_2010': {'offer_px': 27.34, 'premium_pct': 31.6, 'premium_basis': 'FE board: Feb 10, 2010 close; implied price $27.34 at Feb 8 in both opinions (GS 30.9%)'},
 'wisconsinenergy_integrys_2014': {'offer_px': 71.47, 'premium_pct': 17.3, 'premium_basis': '1-day, 6/20/2014 close; 22.8% to 30-day VWAP (Integrys board)'},
 'wps_peoples_2006': {'peoples_exch_px': 41.39, 'premium_basis': 'implied consideration 0.825 x WPS $50.17 (2006-07-05) per proxy'},
 'aes_ipalco_2000': {'premium_pct': 16.3, 'premium_basis': "board: last trading day before announcement (July 14 close $21.50); UBS 18.7% (Jul 12) / 27.0% (1 month)",
                     '_structure_note': 'floating ratio: $25.00 / AES 20-day average, max 0.794 (proxy); 0.463 not printed'},
 'emera_teco_2015': {'premium_pct': 48.3, 'premium_basis': '1-day, unaffected close $18.58 on 2015-07-15 (Moelis; board ~48%)'},
 'exelon_constellation_2011': {'premium_pct': 16.3, 'premium_basis': 'Apr 26, 2011 close (Exelon board, Goldman); 12.5% Apr 27; 20.9% Apr 6 unaffected'},
 'exelon_pepco_2014': {'premium_pct': 29.5, 'premium_basis': '20-day VWAP ending Apr 25, 2014 (proxy); ~24.7% to the $21.85 Apr 25 close (computed)'},
 'fortis_itc_2016': {'premium_pct': 33.0, 'premium_basis': '~33% to 2015-11-27 unaffected close; ~37% to 30-day average (ITC board, at $44.90)',
                     '_structure_note': 'ITC holders own up to ~29% of Fortis post-close (proxy, May 13, 2016)'},
 'greatplains_aquila_2007': {'announced': '2007-02-07', 'premium_pct': -1.3, 'premium_basis': 'discount to Feb 2, 2007 close (GXP board); -0.7% vs Feb 5 (Blackstone)'},
 'iif_sji_2022': {'premium_pct': 51.5, 'premium_basis': '1-day, close 2022-02-22; 47.5% to 30-day VWAP (proxy)'},
 'midamerican_nvenergy_2013': {'announced': '2013-05-29', 'premium_pct': 20.3, 'premium_basis': '1-day, $19.75 close 2013-05-28; 14.4% to 30-day VWAP'},
 'nee_dominion_2026': {'premium_pct': 23.1, 'premium_basis': 'board: over the May 15, 2026 close'},
 'sjw_connecticutwater_2018': {'premium_pct': 33.0, 'premium_basis': "'more than 33%' to $52.57 (proxy)"},
 'southern_agl_2015': {'premium_pct': 37.9, 'premium_basis': '1-day, Aug 21, 2015 close $47.86 (Goldman, board)'},
 'agl_nicor_2010': {'funding': '40% cash / 60% stock', '_structure_note': '33%/67% is the pro forma Nicor/AGL ownership split, not the consideration mix'},
 'firstenergy_gpu_2000': {'_structure_note': 'exchange-ratio collar 1.2318 (min) to 1.5055 (max) around $36.50 of FE stock (proxy)'},
}
# advisor lists the proxy corrects (role conflicts), beyond the additions merged from the extracts
ADV_SET = {
 'iberdrola_energyeast_2007': {'target_advisors': ['JPMorgan', 'Greenhill'], 'acquirer_advisors': ['Banc of America Securities']},
 'macquarie_duquesne_2006': {'target_advisors': ['Lehman Brothers', 'Morgan Stanley']},
}


FAMILY = [('bofa', r'bofa|merrill|banc of america|bank of america'), ('citi', r'citi|salomon'), ('jpm', r'j\.?\s?p\.?\s?morgan|jpmorgan|chase'),
          ('gs', r'goldman'), ('ms', r'morgan stanley'), ('lazard', r'lazard'), ('barclays', r'barclays'), ('ubs', r'\bubs\b'),
          ('blackstone', r'blackstone'), ('cs', r'credit suisse'), ('wf', r'wells fargo'), ('evercore', r'evercore'),
          ('moelis', r'moelis'), ('guggenheim', r'guggenheim'), ('rbc', r'\brbc\b'), ('lehman', r'lehman'), ('greenhill', r'greenhill'),
          ('centerview', r'centerview'), ('scotia', r'scotia'), ('mizuho', r'mizuho'), ('houlihan', r'houlihan')]
SHORT = {'citi': 'Citi', 'lazard': 'Lazard', 'ms': 'Morgan Stanley', 'gs': 'Goldman Sachs', 'blackstone': 'Blackstone', 'ubs': 'UBS'}


def norm_bank(b):
    low = (b or '').lower()
    for k, rx in FAMILY:
        if re.search(rx, low): return k
    return re.sub(r'[^a-z]', '', low)[:8]


def display_bank(b):
    """as-of house name without legal suffixes ('Lazard Freres & Co. LLC' -> 'Lazard', 'Citigroup Global Markets Inc.' -> 'Citi')."""
    k = norm_bank(b)
    if k in SHORT: return SHORT[k]
    b = re.sub(r'\s*\(.*?\)', '', b)
    b = re.sub(r',?\s+(Inc\.?|LLC|L\.P\.|Incorporated|& Co\.?|Co\.)$', '', b.strip())
    return re.sub(r',?\s+(Inc\.?|LLC)$', '', b.strip())


def recompute(d, old_px, new_px):
    r, m = d['raw'], d.setdefault('multiples', {})
    notes = []
    if r.get('eps_fy1'):
        m['pe'] = round(new_px / r['eps_fy1'], 1); notes.append('P/E')
    if r.get('equity_value_usd_b') and old_px:
        eq_old = r['equity_value_usd_b']; eq_new = round(eq_old * new_px / old_px, 3)
        r['equity_value_usd_b'] = eq_new
        if r.get('fv_usd_b'):
            r['fv_usd_b'] = round(r['fv_usd_b'] + (eq_new - eq_old), 3)
            if r.get('ebitda_fy1_usd_m'): m['ev_ebitda'] = round(1000 * r['fv_usd_b'] / r['ebitda_fy1_usd_m'], 1); notes.append('EV/EBITDA')
            rb = r.get('rate_base_usd_b') or (r.get('rate_base_usd_m') or 0) / 1000
            if rb: m['ev_rb'] = round(r['fv_usd_b'] / rb, 1); notes.append('EV/RB')
        notes.append('equity value / EV scaled to the proxy price')
    return notes


def main():
    a = sys.argv[1:]
    rev = a[a.index('--review') + 1]
    prec_p = a[a.index('--precedents') + 1] if '--precedents' in a else os.path.join(BASE, 'data', 'precedents.json')
    mv_p = a[a.index('--valuation') + 1] if '--valuation' in a else os.path.join(BASE, 'data', 'merger_valuation.json')
    doc = json.load(open(prec_p, encoding='utf-8'))
    bak = prec_p + '.bak-preFEreview-20260925'
    if not os.path.exists(bak) and '--no-backup' not in a: shutil.copy(prec_p, bak)
    D = {d['id']: d for d in doc['deals']}
    log = []
    wb = load_workbook(rev, data_only=True)
    # ---- links
    rows = []
    for sh in ('1_Weak_links', '2_Verify_flags'):
        ws = wb[sh]; hdr = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2):
            v = dict(zip(hdr, [c.value for c in r]))
            if v.get('Deal'): rows.append((sh, v))
    # sheet 2 applied after sheet 1 so its corrected URL wins where both touch the same field
    for sh, v in rows:
        d = D.get(v['Deal']); f = FIELD.get(v['Field'], v['Field'])
        if not d: log.append(f'?? unknown deal {v["Deal"]}'); continue
        L = d['links']; dec = (v.get('Decision') or '').strip(); new = (v.get('Corrected URL') or '').strip()
        if (sh, v['Deal'], f) in SKIP:
            log.append(f'SKIPPED {v["Deal"]} {f}: sheet-1 URL is for another deal ({new[:60]}); sheet-2 value applied'); continue
        if dec == 'Confirm':
            L.setdefault('_verified', {})[f] = 'human'; L.setdefault('_link_confidence', {})[f] = 'verified'
            if sh == '2_Verify_flags' and f in (L.get('_link_verify') or {}):
                L['_link_verify'][f] = {'ok': True, 'why': (L['_link_verify'][f].get('why') or '') + f' - confirmed by FE {TODAY}'}
            log.append(f'confirm  {v["Deal"]:34} {f}')
        elif dec == 'Replace' and new:
            L[f] = new; L.setdefault('_verified', {})[f] = 'human'; L.setdefault('_link_confidence', {})[f] = 'verified'
            (L.get('_link_verify') or {}).pop(f, None)
            if v.get('Comment'): L[f'_{f}_note'] = f'FE {TODAY}: {v["Comment"]}'
            log.append(f'replace  {v["Deal"]:34} {f} -> {new}')
        elif dec == 'Remove':
            old = L.get(f); L[f] = None
            (L.get('_link_confidence') or {}).pop(f, None); (L.get('_link_verify') or {}).pop(f, None)
            nf = L.setdefault('_needs_find', []); nf.append(f) if f not in nf else None
            L[f'_{f}_removed'] = f'FE {TODAY}: removed ({old})'
            log.append(f'remove   {v["Deal"]:34} {f}')
        elif v.get('Comment'):
            log.append(f'COMMENT  {v["Deal"]:34} {f}: {v["Comment"]}')
    # Duquesne: the 'deck' is the Aug 9, 2006 Q2 earnings-call transcript (8-K ex-99.2) - keep it as the transcript
    dq = D.get('macquarie_duquesne_2006')
    if dq and dq['links'].get('deck', '').endswith('l21807aexv99w2.htm'):
        L = dq['links']; L['transcript'] = L['deck']; L['_transcript_src'] = 'Q2 2006 earnings call transcript, Aug 9, 2006 (8-K ex-99.2, filed 2006-08-11) - discusses the deal; FE 2026-09-25'
        L['deck'] = None; L.setdefault('_needs_find', []).append('deck'); L.setdefault('_needs_find', []).append('press_release')
        (L.get('_link_confidence') or {}).pop('deck', None)
        log.append('moved    macquarie_duquesne_2006 deck -> transcript; announcement deck / press release still to find')
    # ---- stale flags
    for d in D.values():
        L = d['links']
        for f, s in list((L.get('_link_confidence') or {}).items()):
            if s in ('weak', 'suspect') and not L.get(f):
                L['_link_confidence'].pop(f); log.append(f'stale    {d["id"]:34} {f} (flag dropped)')
    # ---- proxy facts
    mv = json.load(open(mv_p, encoding='utf-8'))['deals'] if os.path.exists(mv_p) else {}
    for did, ed in EDITS.items():
        d = D[did]; r = d['raw']; ch = []
        for k, v in ed.items():
            if k == 'announced':
                ch.append(f'announced {d["announced"]} -> {v}'); d['announced'] = v
            elif k == '_structure_note':
                d['_proxy_structure_note'] = v; ch.append('structure note')
            elif k == 'offer_px':
                old = r.get('offer_px'); r['offer_px'] = v; ch.append(f'offer {old} -> {v}')
                ch += recompute(d, old, v)
            else:
                old = r.get(k); r[k] = v
                if k != 'premium_basis': ch.append(f'{k} {old} -> {v}')
        d['_proxy_review'] = f'FE {TODAY}: proxy value adopted ({"; ".join(ch)}); source {((mv.get(did) or {}).get("document") or {}).get("url")}'
        log.append(f'proxy    {did:34} ' + '; '.join(ch))
    for did, x in mv.items():
        d = D.get(did)
        if not d: continue
        pr = x.get('process') or {}; tf = pr.get('termination_fee') or {}
        d['proxy'] = {'form': (x.get('document') or {}).get('form'), 'filed': (x.get('document') or {}).get('filed'),
                      'url': (x.get('document') or {}).get('url'),
                      'termination_fee_target_m': tf.get('target_pays_m'), 'reverse_fee_m': tf.get('reverse_fee_m'),
                      'termination_fee_pct_equity': tf.get('pct_equity_value'),
                      'advisors': [{'bank': b.get('bank'), 'side': b.get('side'), 'opinion': bool(b.get('methods'))} for b in x.get('advisors') or []],
                      'process_type': pr.get('type'), 'bidders_final': pr.get('final_bids')}
        L = d['links']
        if not L.get('fairness_opinion') and d['proxy']['url']:
            L['fairness_opinion'] = d['proxy']['url']
            L['_fo_src'] = f'merger proxy {d["proxy"]["form"]} {d["proxy"]["filed"] or ""} (fetch_merger_proxies.py; valuation extracted to merger_valuation.json)'
            log.append(f'fo link  {did:34} -> {d["proxy"]["form"]}')
        adv = d.setdefault('advisors', {'target_advisors': [], 'acquirer_advisors': []})
        if did in ADV_SET:
            for k, v in ADV_SET[did].items():
                adv[k] = v
            log.append(f'advisors {did:34} set from proxy {ADV_SET[did]}')
        else:
            for b in x.get('advisors') or []:
                side = {'target': 'target_advisors', 'special_committee': 'target_advisors', 'acquirer': 'acquirer_advisors'}.get(b.get('side'))
                if not side or not b.get('bank'): continue
                lst = adv.setdefault(side, [])
                if not any(norm_bank(b['bank']) == norm_bank(e) for e in lst):
                    nm = display_bank(b['bank']); lst.append(nm); log.append(f'advisors {did:34} + {nm} ({b["side"]}) from proxy')
        named = {norm_bank(b.get('bank') or '') for b in x.get('advisors') or []}
        missing = [e for k in ('target_advisors', 'acquirer_advisors') for e in adv.get(k, []) if norm_bank(e) not in named]
        if missing:
            d['_advisor_note_proxy'] = f'not named in the {d["proxy"]["form"]} (kept from the deal file; the other side\'s filings may name them): ' + ', '.join(missing)
    doc.setdefault('_review_log', []).append({'date': TODAY, 'source': os.path.basename(rev), 'changes': len(log)})
    json.dump(doc, open(prec_p, 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
    print('\n'.join(log)); print(f'\n{len(log)} changes -> {prec_p}')


if __name__ == '__main__':
    main()
